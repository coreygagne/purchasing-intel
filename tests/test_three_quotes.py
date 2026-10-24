"""Distinct supplier comparisons must use delivered costs when they are known."""

from app.db import Database
from app.three_quotes import workflow


def _offer(oid, price, shipping):
    components = [{"label": "Goods (CAD)", "value": price * 15, "basis": "calculated"}]
    missing = []
    if shipping is None:
        missing.append("shipping")
    else:
        components.append({"label": "Shipping", "value": shipping, "basis": "quoted"})
    return {"id": oid, "comparable": True, "price_cad": price, "price": price, "currency": "CAD",
            "price_basis": "structured_data", "confidence": "high", "price_kind": "asking",
            "price_type": "listed", "compare_value": price, "price_label": "seller page",
            "url": f"https://{oid}.example/item", "observed_at": "2026-10-23T12:00:00Z",
            "delivered": {"value": price * 15 + (shipping or 0), "complete": not missing,
                          "missing": missing, "components": components}}


def _vendor(name, domain, ids):
    return {"id": domain, "name": name, "domain": domain, "type": "retailer", "business_kind": "seller",
            "offer_ids": ids, "emails": [], "phones": [], "website": f"https://{domain}/",
            "tier": "canada", "lists_item": True, "lowest": {"comparable": True}}


def _result():
    return {"quantity": 15, "identity": {"model_tokens": ["PX-123"]},
            "offers": [_offer("a", 9, None), _offer("a2", 11, 20),
                       _offer("b", 10, 25), _offer("c", 12, 0)],
            "vendors": [_vendor("A", "a.example", ["a", "a2"]),
                        _vendor("A duplicate", "a.example", ["a"]),
                        _vendor("B", "b.example", ["b"]), _vendor("C", "c.example", ["c"])],
            "rfq_drafts": []}


def test_three_options_have_distinct_suppliers_and_shipping():
    result = workflow(_result(), [], {}, None, {})
    options = result["best_options"]
    assert len(options) == 3
    assert len({o["supplier_key"] for o in options}) == 3
    assert [o["comparison_value_cad"] for o in options] == [175, 180, 185]
    assert all(o["delivered_ready"] for o in options)


def test_unknown_shipping_is_not_ranked_as_cheapest_delivered():
    data = _result()
    data["offers"] = [o for o in data["offers"] if o["id"] != "a2"]
    result = workflow(data, [], {}, None, {})
    assert result["best_options"][-1]["supplier_name"] == "A"
    assert result["best_options"][-1]["comparison_value_cad"] is None
    assert result["delivered_ready_count"] == 2
    assert result["three_prices_found"] is True


def test_received_reply_replaces_own_published_price_and_requires_shipping():
    quote = {"supplier_key": "domain:a.example", "supplier_name": "A", "status": "received",
             "quantity": 15, "unit_price": 8, "currency": "CAD", "fx_rate_to_cad": None,
             "shipping_cad": 12, "other_cad": 0, "specs_confirmed": True, "valid_until": None}
    result = workflow(_result(), [quote], {}, None, {})
    assert len({o["supplier_key"] for o in result["best_options"]}) == 3
    assert next(o for o in result["best_options"] if o["supplier_key"] == quote["supplier_key"])["kind"] == "supplier_quote"
    quote["shipping_cad"] = None
    result = workflow(_result(), [quote], {}, None, {})
    assert result["comparable_count"] == 0
    assert "shipping cost or pickup not confirmed" in result["quotes"][0]["issues"]


def test_quote_and_backup_survive_restore(tmp_path):
    db = Database(tmp_path / "live.db")
    cid = db.save_comparison(title="pylons", query="pylons", normalized="pylons", product_key="pylons",
                             result=_result())
    row = {"supplier_key": "domain:a.example", "supplier_name": "A", "method": "phone", "status": "received",
           "email": None, "phone": "555-0100", "quantity": 15, "unit_price": 8.0, "currency": "CAD",
           "fx_rate_to_cad": None, "shipping_cad": 12.0, "other_cad": 0, "lead_days": 3,
           "valid_until": None, "specs_confirmed": True, "reference": "Q-1", "notes": ""}
    db.put_quote(cid, row)
    backup = db.backup_bytes()
    row["unit_price"] = 20.0
    db.put_quote(cid, row)
    before = db.restore_bytes(backup)
    assert before.exists()
    assert db.quotes_for(cid)[0]["unit_price"] == 8.0
    db.close()


def test_quote_api_and_backup_download(tmp_path):
    from fastapi.testclient import TestClient
    from app.config import Settings
    from app.main import create_app
    from app.research import ANALYSIS_VERSION

    app = create_app(Settings(database_path=tmp_path / "api.db"))
    result = _result() | {"analysis_version": ANALYSIS_VERSION}
    cid = app.state.db.save_comparison(title="pylons", query="pylons", normalized="pylons",
                                       product_key="pylons", result=result)
    with TestClient(app, base_url="http://127.0.0.1:8000", client=("127.0.0.1", 50123)) as client:
        response = client.get(f"/api/comparisons/{cid}")
        assert response.status_code == 200
        assert len(response.json()["quote_workflow"]["best_options"]) == 3
        profile = {"supplier_key": "domain:a.example", "name": "A", "preferred": True,
                   "email": "sales@a.example", "phone": "555-0100"}
        assert client.post("/api/supplier-profiles", json=profile).status_code == 200
        quote = {"supplier_key": "domain:a.example", "supplier_name": "A", "method": "phone",
                 "status": "received", "quantity": 15, "unit_price": 8, "currency": "CAD",
                 "shipping_cad": 12, "specs_confirmed": True}
        assert client.put(f"/api/comparisons/{cid}/quotes", json=quote).status_code == 200
        updated = client.get(f"/api/comparisons/{cid}").json()["quote_workflow"]
        assert updated["comparable_count"] == 1
        assert updated["profiles"]["domain:a.example"]["preferred"] is True
        assert len({o["supplier_key"] for o in updated["best_options"]}) == 3
        backup = client.get("/api/backup.db")
        assert backup.status_code == 200 and backup.content.startswith(b"SQLite format 3\x00")


def test_excel_contains_distinct_options_and_quote_tracker(tmp_path):
    from io import BytesIO
    from openpyxl import load_workbook
    from app.export import comparison_xlsx
    from tests.test_discovery import local_web, run

    result = run(tmp_path, local_web(), query="Acme TX-4400 torque widget", quantity=2)
    flow = workflow(result, [], {}, None, {})
    data = comparison_xlsx({"result": result, "quote_workflow": flow}, [])
    wb = load_workbook(BytesIO(data), read_only=True)
    assert "Three suppliers" in wb.sheetnames and "Quote tracker" in wb.sheetnames
    suppliers = [r[0] for r in list(wb["Three suppliers"].values)[1:]]
    assert len(suppliers) == len(set(suppliers))
