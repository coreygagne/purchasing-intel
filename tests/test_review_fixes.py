"""Regression tests for the findings of the Phase 1 code review."""

import io
import json
from types import SimpleNamespace

import httpx
import pytest
from openpyxl import load_workbook

from app import extract as ex
from app.config import Settings
from app.db import Database
from app.evidence import OfferBuilder, Source, build_sources
from app.export import comparison_xlsx, offers_csv
from app.fetch import PageFetcher
from app.llm import ClaudeAnalyst
from app.main import price_history
from app.providers.base import SearchHit
from app.research import _brand_words, _clean_title
from tests.fakeweb import PRODUCT, FakeWeb, jsonld_page, tavily_result
from tests.test_api import make_client, research


def result_of(c, job):
    assert job["status"] == "done", job
    return c.get(f"/api/comparisons/{job['comparison_id']}").json()["result"]


# --- C1: an aftermarket page used to crash the whole research with KeyError ------------------------
def test_compatible_product_page_does_not_crash(tmp_path):
    fw = FakeWeb()
    fw.pages["https://www.kmstools.com/acme-tx-4400-torque-widget"] = (
        200, jsonld_page("Replacement anvil for Acme TX-4400", "39.99", "CAD"))
    with make_client(tmp_path, fw) as c:
        res = result_of(c, research(c))
        kms = [o for o in res["offers"] if o["supplier"] == "KMS Tools" and o["price_basis"] == "structured_data"]
        assert kms and kms[0]["match_level"] == "compatible" and not kms[0]["comparable"]


# --- C2: the AI may only contribute verified page facts --------------------------------------------
def _src(sid, text, *, title="Acme TX-4400 Torque Widget", type_="retailer", purposes=("price",)):
    return Source(id=sid, url=f"https://shop{sid}.ca/x", domain=f"shop{sid}.ca", title=title, snippet="",
                  type=type_, name=f"Shop {sid}", country="CA", country_basis="domain suffix", in_directory=False,
                  purposes=list(purposes), queries=[], provider="Tavily", retrieved_at="2026-10-20T00:00:00+00:00",
                  page_text=text)


def _analyst(payload):
    async def create(**kw):
        return SimpleNamespace(stop_reason="end_turn", model="m",
                               content=[SimpleNamespace(type="text", text=json.dumps(payload))],
                               usage=SimpleNamespace(input_tokens=1, output_tokens=1))
    msgs = SimpleNamespace(create=create)
    return ClaudeAnalyst(Settings(anthropic_api_key="k"), client=SimpleNamespace(messages=msgs, beta=SimpleNamespace(messages=msgs)))


def _offer(sid, price, quote, **kw):
    return {"source_id": sid, "product_title": kw.get("title", "Acme TX-4400"), "supplier_name": kw.get("supplier", ""),
            "price_text": price, "currency_stated": "", "pack_quantity": kw.get("pack", 0),
            "condition": kw.get("condition", "not stated"), "evidence_quote": quote, "shipping_text": "",
            "warranty": "", "moq": "", "availability": ""}


async def test_ai_offers_take_identity_pack_condition_and_seller_from_the_source():
    ident = SimpleNamespace(name="Acme TX-4400", model_tokens=["TX-4400"], key_specs=[])
    accessory_page = ("Acme TX-4400 Torque Widget. Specifications and details for the widget. Lots of text here "
                      "that goes on for a while about the widget design and its features and uses. Customers also "
                      "bought: Spare Battery Pack $199.00 each.")
    sources = [
        _src("S1", accessory_page),
        _src("S2", "Acme TX-4400 Torque Widget, our price $310.00 each."),
        _src("S3", "Acme TX-4400 Torque Widget. Our price $250.00", title="Acme TX-4400 (Factory Reconditioned)"),
        _src("S4", "We tested the Acme TX-4400 and at $249.00 it is a bargain.", type_="professional review",
             purposes=("reviews",)),
    ]
    payload = {"offers": [
        _offer("S1", "$199.00", "Spare Battery Pack $199.00 each", supplier="Home Depot"),   # accessory
        _offer("S2", "$310.00", "Acme TX-4400 Torque Widget, our price $310.00 each", pack=20, condition="new",
               title="Acme TX-4400 (20 pack)"),                                                  # invented pack
        _offer("S3", "$250.00", "Acme TX-4400 Torque Widget. Our price $250.00", condition="new"),  # refurb title
        _offer("S4", "$249.00", "at $249.00 it is a bargain", supplier="Home Depot"),           # review page
    ], "quality": {"praise": [], "criticism": [], "reliability_issues": [], "overall_note": ""}, "alternatives": []}
    builder = OfferBuilder(["Acme"])
    out = await _analyst(payload).extract(ident, sources, builder)
    reasons = out["verification"]["offer_rejection_reasons"]
    assert out["verification"]["offers_rejected"] == 2                       # accessory + review page
    assert any("not a seller page" in r for r in reasons)
    by_src = {o.source_id: o for o in builder.offers}
    assert set(by_src) == {"S2", "S3"}
    assert by_src["S2"].pack_qty is None and by_src["S2"].supplier == "Shop S2"
    assert by_src["S2"].product_title == "Acme TX-4400 Torque Widget"       # from the source, not the AI
    assert by_src["S3"].condition == "refurbished"                           # source title wins over AI "new"


# --- C3: stray stock/related-product text must not become a pack size ------------------------------
def test_stock_text_is_not_a_pack_size_and_mixed_packs_keep_all_offers(tmp_path):
    fw = FakeWeb()
    fw.tavily_results["price"][1]["content"] = ("Acme TX-4400 torque widget $349.99 Only 3 units left. Over 500 units "
                                                "sold. Customers also bought: battery 2-pack.")
    fw.pages["https://www.kmstools.com/acme-tx-4400-torque-widget"] = (404, "gone")
    with make_client(tmp_path, fw) as c:
        res = result_of(c, research(c))
        kms = next(o for o in res["offers"] if o["supplier"] == "KMS Tools")
        assert kms["pack_qty"] is None
    pool = ex.pack_quantity("Acme TX-4400 Torque Widget (2-Pack)")
    assert pool and pool[0] == 2


def test_mixed_pack_sizes_do_not_drop_single_listings():
    from app.analysis import recommend, set_compare_values
    from app.evidence import Offer

    def mk(oid, price, pack=None):
        o = Offer(id=oid, source_id=None, url=f"https://s{oid}.ca", supplier=f"S{oid}", supplier_domain="x.ca",
                  supplier_type="retailer", supplier_in_directory=True, country="CA", country_basis="directory",
                  product_title="X", price=price, price_high=None, currency="CAD", currency_basis="stated",
                  price_text=str(price), price_type="listed", price_basis="structured_data", evidence_quote="",
                  observed_at="2026-10-20", provider="t", match_level="exact", pack_qty=pack, confidence="high")
        o.price_cad = price
        o.unit_price_cad = price / pack if pack else None
        o.comparable = True
        return o
    # one listing states a pack size: compare listing prices; nothing is dropped
    offers = [mk("1", 100.0), mk("2", 110.0), mk("3", 379.0, pack=3)]
    set_compare_values(offers)
    rec = recommend(offers, ["X-1000"])
    assert rec["lowest"]["offer_id"] == "1" and rec["pool_size"] == 3 and rec["basis"] == "listing price (CAD)"
    # two listings state pack sizes: per-unit, singles counted as one unit and flagged
    offers = [mk("1", 100.0), mk("2", 250.0, pack=3), mk("3", 379.0, pack=5)]
    set_compare_values(offers)
    rec = recommend(offers, ["X-1000"])
    assert rec["lowest"]["offer_id"] == "3" and "counted as single units" in rec["basis"]
    assert any("counted as a single unit" in f for f in offers[0].flags)


# --- C4: several offers under one Product -> unclear, not ranked ------------------------------------
def test_multi_offer_page_is_unclear_not_ranked(tmp_path):
    fw = FakeWeb()
    page = {"@context": "https://schema.org", "@type": "Product", "name": "Acme TX-4400 Torque Widget", "mpn": "TX-4400",
            "offers": [{"@type": "Offer", "price": "199.00", "priceCurrency": "CAD", "sku": "TX-4400-REFURB-NOBOX"},
                       {"@type": "Offer", "price": "339.00", "priceCurrency": "CAD", "sku": "TX-4400"}]}
    fw.pages["https://www.kmstools.com/acme-tx-4400-torque-widget"] = (
        200, f'<script type="application/ld+json">{json.dumps(page)}</script>')
    with make_client(tmp_path, fw) as c:
        res = result_of(c, research(c))
        kms = [o for o in res["offers"] if o["supplier"] == "KMS Tools"]
        assert len(kms) == 1 and kms[0]["price_type"] == "unclear" and not kms[0]["comparable"]
        assert res["recommendations"]["lowest"]["supplier"] != "KMS Tools"


# --- I5 / I6: spec numbers and genuine "compatible with" titles -------------------------------------
def test_spec_numbers_are_not_model_numbers():
    assert ex.model_tokens("1500 watt space heater") == []
    assert ex.model_tokens("Impact wrench 1/2 inch 1000 ft-lb") == []
    assert ex.classify_input("1500 watt space heater") == "text"


def test_genuine_listings_not_marked_aftermarket_or_used():
    assert ex.match_offer(["2767-20"], "Milwaukee 2767-20 M18 FUEL - Compatible with M18 Batteries").level == "exact"
    assert ex.match_offer(["2767-20"], "2767-20 fits 1/2 in. sockets").level == "exact"
    assert ex.match_offer(["2767-20"], "Replacement anvil for Milwaukee 2767-20").level == "compatible"
    assert ex.detect_condition("Wrench is used by mechanics everywhere") is None
    assert ex.detect_condition("Milwaukee 2767-20 (Used)") == "used"


# --- I7: URLs that differ by query string are different listings ----------------------------------
def test_query_string_urls_stay_separate():
    hits = [SearchHit("T", "q", "price", "https://shop.ca/ProductDetail.aspx?sku=2767-22", "Kit 2767-22", "$429", 1),
            SearchHit("T", "q", "price", "https://shop.ca/ProductDetail.aspx?sku=2767-20&utm_source=x", "Bare 2767-20", "$289", 2),
            SearchHit("T", "q", "reviews", "https://shop.ca/ProductDetail.aspx?sku=2767-20", "Bare 2767-20", "$289", 3)]
    srcs = build_sources(hits, [])
    assert len(srcs) == 2 and srcs[1].purposes == ["price", "reviews"]


# --- I8 / I9: decimal commas and "from" prices -----------------------------------------------------
def test_decimal_comma_prices():
    assert ex.parse_price_text("349,99").amount == 349.99
    assert ex.parse_price_text("1.299,99").amount == 1299.99
    assert [p.amount for p in ex.find_prices("€1.299,99")] == [1299.99]
    assert ex.find_prices("$3.4567") == []  # unreadable precision: not guessed


def test_aggregate_offer_low_price_is_a_from_price(tmp_path):
    fw = FakeWeb()
    page = {"@type": "Product", "name": "Acme TX-4400 Torque Widget",
            "offers": {"@type": "AggregateOffer", "lowPrice": "149.00", "priceCurrency": "CAD", "offerCount": 12}}
    fw.pages["https://www.kmstools.com/acme-tx-4400-torque-widget"] = (
        200, f'<script type="application/ld+json">{json.dumps(page)}</script>')
    with make_client(tmp_path, fw) as c:
        res = result_of(c, research(c))
        kms = next(o for o in res["offers"] if o["supplier"] == "KMS Tools" and o["price_basis"] == "structured_data")
        assert kms["price_type"] == "range" and not kms["comparable"]


# --- I10: AI quality summary may not introduce statistics -------------------------------------------
async def test_ai_quality_points_cannot_add_statistics():
    ident = SimpleNamespace(name="Acme TX-4400", model_tokens=["TX-4400"], key_specs=[])
    src = _src("S1", "Owners say the Acme TX-4400 is reliable and well built for daily shop use.",
               type_="forum / user discussion", purposes=("reviews",))
    payload = {"offers": [], "alternatives": [], "quality": {
        "praise": [{"point": "92% of owners rate it 5 stars", "source_ids": ["S1"],
                    "quotes": ["the Acme TX-4400 is reliable and well built"]},
                   {"point": "Owners find it reliable", "source_ids": ["S1"], "quotes": ["4400"]},
                   {"point": "Owners call it reliable and well built", "source_ids": ["S1"],
                    "quotes": ["the Acme TX-4400 is reliable and well built"]}],
        "criticism": [], "reliability_issues": [], "overall_note": "Rated 4.8/5 across 3,000 reviews."}}
    out = await _analyst(payload).extract(ident, [src], OfferBuilder([]))
    q = out["quality"]
    assert [p["point"] for p in q["praise"]] == ["Owners call it reliable and well built"]
    assert q["overall_note"] == "" and q["verification"]["points_rejected"] == 2


# --- I11: SSRF guard ---------------------------------------------------------------------------------
@pytest.mark.parametrize("url,expect", [
    ("http://127.0.0.1/admin", "private or internal"),
    ("http://[::1]/", "private or internal"),
    ("http://169.254.169.254/latest/meta-data/", "private or internal"),
    ("http://intranet.internal.test/", "private or internal"),
    ("http://shop.example.ca:8080/", "port 8080"),
    ("ftp://shop.example.ca/file", "only http"),
    ("http://user:pw@shop.example.ca/", "credentials"),
])
async def test_fetcher_blocks_internal_and_odd_urls(url, expect):
    calls = []
    client = httpx.AsyncClient(transport=httpx.MockTransport(lambda r: calls.append(r) or httpx.Response(200, text="x")))
    f = PageFetcher(client, user_agent="t", timeout_s=2, max_bytes=1000, respect_robots=False)
    r = await f.fetch(url)
    assert not r.ok and expect in r.error and calls == []


async def test_fetcher_blocks_redirect_into_private_network():
    def handler(request):
        if request.url.host == "shop.example.ca":
            return httpx.Response(302, headers={"location": "http://10.0.0.7/secret"})
        raise AssertionError("must not reach the private address")
    f = PageFetcher(httpx.AsyncClient(transport=httpx.MockTransport(handler)), user_agent="t", timeout_s=2,
                    max_bytes=1000, respect_robots=False)
    r = await f.fetch("https://shop.example.ca/p")
    assert not r.ok and "private or internal" in r.error


def test_private_url_input_is_not_fetched(tmp_path):
    fw = FakeWeb()
    with make_client(tmp_path, fw) as c:
        res = result_of(c, research(c, "http://192.168.1.10/products/acme-tx-4400-torque-widget"))
        assert not any("192.168.1.10" in call for call in fw.calls)
        assert any("private or internal" in g for g in res["gaps"])


# --- I12: spreadsheet exports are inert and robust --------------------------------------------------
def test_exports_escape_formulas_and_strip_control_chars(tmp_path):
    fw = FakeWeb()
    fw.tavily_results["price"][2]["title"] = '=HYPERLINK("http://evil.example","Acme TX-4400")\x0b'
    with make_client(tmp_path, fw) as c:
        cid = research(c)["comparison_id"]
        csv_text = c.get(f"/api/comparisons/{cid}/export.csv").text
        assert "'=HYPERLINK" in csv_text and "\x0b" not in csv_text
        r = c.get(f"/api/comparisons/{cid}/export.xlsx")
        assert r.status_code == 200
        wb = load_workbook(io.BytesIO(r.content))
        cells = [cell for row in wb["Offers"].iter_rows() for cell in row
                 if isinstance(cell.value, str) and "HYPERLINK" in cell.value]
        assert cells and all(cell.data_type == "s" and cell.value.startswith("'") for cell in cells)


# --- I14: no-model queries are never charted as one trend -------------------------------------------
def test_history_does_not_chart_unverified_matches():
    obs = [{"observed_at": f"2026-10-0{d}T00:00:00", "price_cad": 50.0 + d, "price": 50.0 + d, "currency": "CAD",
            "pack_qty": None, "url": f"https://s{d}.ca", "price_basis": "structured_data", "supplier": f"S{d}",
            "supplier_domain": "s.ca", "match_level": "unverified", "condition": None, "price_type": "listed"}
           for d in (1, 2, 3, 4)]
    h = price_history(obs)
    assert h["chartable"] is False and "No model number" in h["reason"]


# --- I15: description words are not brands ----------------------------------------------------------
def test_description_word_does_not_make_a_manufacturer(tmp_path):
    assert _brand_words("Nitrile gloves powder free") == []
    fw = FakeWeb(tavily_results={"price": [tavily_result("https://nitriledirect-deals.com/gloves", "Nitrile gloves 100 count",
                                                         "Nitrile gloves 100 count $12.99")]})
    with make_client(tmp_path, fw) as c:
        res = result_of(c, research(c, "nitrile gloves powder free"))
        src = next(s for s in res["sources"] if s["domain"] == "nitriledirect-deals.com")
        assert src["type"] != "manufacturer"


# --- Minor items --------------------------------------------------------------------------------------
def test_out_of_stock_and_moq_over_quantity_are_not_comparable(tmp_path):
    fw = FakeWeb()
    fw.pages["https://www.kmstools.com/acme-tx-4400-torque-widget"] = (
        200, jsonld_page("Acme TX-4400 Torque Widget", "339.99", "CAD", availability="https://schema.org/OutOfStock"))
    with make_client(tmp_path, fw) as c:
        res = result_of(c, research(c))
        kms = next(o for o in res["offers"] if o["supplier"] == "KMS Tools" and o["price_basis"] == "structured_data")
        assert not kms["comparable"] and any("not available" in r for r in kms["exclusion_reasons"])
        assert res["recommendations"]["best_value"]["supplier"] != "KMS Tools"


def test_refresh_keeps_saved_title(tmp_path):
    with make_client(tmp_path) as c:
        cid = research(c)["comparison_id"]
        c.patch(f"/api/comparisons/{cid}", json={"saved": True, "title": "Shop widget"})
        research(c, comparison_id=cid)
        comp = c.get(f"/api/comparisons/{cid}").json()
        assert comp["title"] == "Shop widget" and comp["saved"] is True


def test_history_returns_newest_when_limited():
    db = Database(":memory:")
    db.add_observations([{"product_key": "k", "product_title": "t", "url": "u", "price_basis": "structured_data",
                          "price_type": "listed", "observed_at": f"2026-10-{d:02d}T00:00:00", "price": float(d)}
                         for d in range(1, 21)])
    rows = db.observations(product_key="k", limit=5)
    assert [r["price"] for r in rows] == [16.0, 17.0, 18.0, 19.0, 20.0]


def test_title_cleaning_keeps_model_numbers():
    assert _clean_title("Acme | TX-4400 widget") == "Acme TX-4400 widget"
    assert _clean_title("Acme TX-4400 Torque Widget | Grainger Canada") == "Acme TX-4400 Torque Widget"


# --- Re-review: dense page text next to the model name -----------------------------------------------
@pytest.mark.parametrize("page,price,quote,accepted", [
    # the reviewer's original reproduction: accessory after the product price
    ("Milwaukee 2767-20 M18 FUEL 1/2 in. impact wrench. Our price $379.00. Frequently bought together: "
     "M18 REDLITHIUM 5.0Ah Battery $199.00", "$199.00", "M18 REDLITHIUM 5.0Ah Battery $199.00", False),
    # another price between the model and the quoted price (no other warning words)
    ("Milwaukee 2767-20 our price $379.00 | M18 Battery $199.00", "$199.00", "M18 Battery $199.00", False),
    # round-3 cases: add-on and "sold separately" wording right after the model
    ("Milwaukee 2767-20 Impact Wrench. Add a M18 battery: $199.00", "$199.00", "Add a M18 battery: $199.00", False),
    ("Milwaukee 2767-20 bare tool, batteries sold separately. Kit price $599.00", "$599.00", "Kit price $599.00", False),
    # sibling kit named between the model and the price
    ("Milwaukee 2767-20 bare tool. Also available: 2767-22 kit with batteries $599.00", "$599.00",
     "2767-22 kit with batteries $599.00", False),
    # related-products wording without another price in between
    ("Milwaukee 2767-20 impact wrench. You may also like: M18 Battery $199.00", "$199.00",
     "M18 Battery $199.00", False),
    # the genuine price right after the model name
    ("Milwaukee 2767-20 M18 FUEL impact wrench, our price $349.00 each", "$349.00", "our price $349.00 each", True),
    # accessory "for" the model
    ("Spare battery for Milwaukee 2767-20 $199.00", "$199.00", "Spare battery for Milwaukee 2767-20 $199.00", True),
])
async def test_ai_price_must_belong_to_the_requested_model(page, price, quote, accepted):
    ident = SimpleNamespace(name="Milwaukee 2767-20", model_tokens=["2767-20"], key_specs=[])
    src = _src("S1", page, title="Milwaukee 2767-20 M18 FUEL Impact Wrench (2-Pack of bits)")
    payload = {"offers": [_offer("S1", price, quote)], "alternatives": [],
               "quality": {"praise": [], "criticism": [], "reliability_issues": [], "overall_note": ""}}
    builder = OfferBuilder(["Milwaukee"])
    await _analyst(payload).extract(ident, [src], builder)
    assert bool(builder.offers) == accepted
    if accepted:
        from app.evidence import apply_match
        apply_match(builder.offers, ["2767-20"], {"S1": src})
        o = builder.offers[0]
        assert o.pack_qty is None or o.pack_text in src.title          # pack never from nearby text
        if "Spare battery" in page:
            assert o.match_level == "compatible"                         # accessory for the model, not the model
        else:
            assert o.match_level == "exact"


def test_bare_for_does_not_mark_genuine_listings_aftermarket():
    for t in ("Best price for 2767-20 in Canada", "Shop now for Milwaukee 2767-20",
              "Looking for a Milwaukee 2767-20? $349.00"):
        assert ex.match_offer(["2767-20"], t).level == "exact", t
    assert ex.match_offer(["2767-20"], "M18 batteries for Milwaukee 2767-20").level == "compatible"


def test_set_piece_counts_are_not_pack_sizes():
    assert ex.pack_quantity("Socket set (28 pcs)") is None
    assert ex.pack_quantity("Nitrile gloves 100 count")[0] == 100
    assert ex.pack_quantity("First aid kit refill, 100 count")[0] == 100
    assert ex.model_tokens("3M 8210 pack of 20") == ["8210"]
    assert ex.model_tokens("Makita 9557 in stock") == ["9557"]


def test_slow_ai_is_skipped_not_fatal(tmp_path, monkeypatch):
    import asyncio as aio

    from app import research as rs
    from app.main import create_app
    from fastapi.testclient import TestClient

    class SlowAI:
        async def plan(self, *a):
            await aio.sleep(5)

        async def extract(self, *a):
            await aio.sleep(5)

    monkeypatch.setattr(rs, "AI_PLAN_TIMEOUT_S", 0.05)
    monkeypatch.setattr(rs, "AI_EXTRACT_TIMEOUT_S", 0.05)
    settings = Settings(tavily_api_key="tvly-test-key", anthropic_api_key="k", database_path=tmp_path / "s.db")
    with TestClient(create_app(settings, transport=FakeWeb().transport(), llm=SlowAI())) as c:
        res = result_of(c, research(c))
        assert res["recommendations"]["lowest"] is not None
        assert sum("took too long and was skipped" in g for g in res["gaps"]) == 2
